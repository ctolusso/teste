# ============================================================
# TYTON – CLIPPING DE NOTÍCIAS (D0)
# Google News + Excel + Email (ANTI-REPETIÇÃO)
# Roda a cada 2 horas e envia só as notícias novas do dia.
# ============================================================

import html as html_lib
import json
import os
import base64
import re
import sys
import unicodedata
import urllib.parse
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

import feedparser
import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# ============================================================
# CONFIGURAÇÕES
# ============================================================

EMAIL_REMETENTE = os.environ.get("CLIPPING_EMAIL_REMETENTE", "tyton@tytoncapital.com.br")
EMAIL_DESTINO = os.environ.get("CLIPPING_EMAIL_DESTINO", "asset@tytoncapital.com.br")

# Envio pelo Microsoft Graph (HTTPS). As credenciais NUNCA ficam no código:
# vêm das variáveis de ambiente, criadas pela TI no Azure (app com permissão Mail.Send)
GRAPH_TENANT_ID = os.environ.get("CLIPPING_TENANT_ID")
GRAPH_CLIENT_ID = os.environ.get("CLIPPING_CLIENT_ID")
GRAPH_CLIENT_SECRET = os.environ.get("CLIPPING_CLIENT_SECRET")

FUSO = ZoneInfo("America/Sao_Paulo")

PASTA_DADOS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dados")
CAMINHO_EXCEL = os.path.join(PASTA_DADOS, "Noticias_Diarias.xlsx")
HISTORICO_JSON = os.path.join(PASTA_DADOS, "historico_clipping.json")

# Sites financeiros confiáveis para query secundária
SITES_FINANCEIROS = (
    "site:infomoney.com.br OR site:valor.com.br OR site:exame.com "
    "OR site:cnnbrasil.com.br OR site:estadao.com.br "
    "OR site:bloomberglinea.com.br OR site:moneytimes.com.br "
    "OR site:braziljournal.com"
)

# ============================================================
# EMISSORES + TERMOS
# ============================================================

EMISSORES = {
    "CMIN3": ["CMIN3", "CSN Mineração", "CSNAA2", "CSNAA1"],
    "DASA3": ["DASA3", "DASA", "Diagnósticos da América", "DASAA5", "DASAA6"],
    "BRAV3": ["BRAV3", "Brava Energia", "Brava Energia SA"],
    "MEAL3": ["MEAL3", "IMC Alimentos", "International Meal Company"],
    "BEEF3": ["BEEF3", "Minerva", "Minerva Foods", "MRNVA1"],
    "MOVI3": ["MOVI3", "Movida Rent a Car", "Movida Aluguel de Carros", "Movida ON", "MVLV19", "MVLVA1", "Movida SA"],
    "OPCT3": ["OPCT3", "OceanPact", "OPCT17", "Ocean Pact"],
    "PRIO3": ["PRIO3", "Petro Rio", "PetroRio", "PFOR16"],
    "PRNR3": ["PRNR3", "Priner", "PRNR12", "PRNR22"],
    "USIM5": ["USIM5", "Usiminas", "Usinas siderúrgicas de Minas Gerais", "USIMA0"],
    "VAMO3": ["VAMO3", "Vamos Locação de Caminhões", "Grupo Vamos", "Vamos Locação", "VAMO13", "VAMO14", "VAMO23", "VAMO24"],
    "WIZC3": ["WIZC3", "Wiz Co Participações", "Wiz Co", "Wiz Soluções"],
    "SIMH3": ["SIMH3", "Simpar SA", "Simpar", "SIMH16", "JSMLA5"],
    "ASAI3": ["ASAI3", "Sendas Distribuidora SA", "Sendas Distribuidora"],
    "ITUB4": ["ITUB4", "Itaú Unibanco", "Itaú BBA"],
    "PGMN3": ["PGMN3", "Pague menos", "PGMN1"],
    "SBSP3": ["SBSP3", "Companhia de Saneamento Básico do Estado de São Paulo"],
    "AURA33": ["AURA33", "Aura Minerals", "AUGO", "Aura Minerals Inc"],
    "PCAR3": ["PCAR3", "Grupo Pão de açúcar", "GPA"],
    "DIRR3": ["DIRR3", "Direcional Engenharia ON", "Direcional Engenharia SA"],
    "CURY3": ["CURY3", "Cury Construtora e Incorporadora SA", "Cury Construtora"],
    "RDOR3": ["RDOR3", "Rede D'Or"],
    "SMFT3": ["SMFT3", "SMFT9", "Smart Fit", "Smartfit Escola de Ginastica e Danca SA"],
    "EQTL3": ["EQTL3", "EQLT3", "Equatorial Energia", "Equatorial SA", "Equatorial ON", "Equatorial Energia SA"],
    "BPAC11": ["BPAC11", "BPAC3"],
    "TFCO4": ["Track & Field SA Co", "TFCO4", "TRACK E FIELD", "Track & Field PN"],
    "JBSS3": ["JBSS3", "JBS ON", "JBS S.A", "JBS"],
    "PSSA3": ["PSSA3", "Porto Seguro ON", "Porto Seguro Itaú Unibanco Participações S.A.", "Porto Seguro S.A", "Porto Seguro SA"],
    "ALOS3": ["ALOS3", "Allos SA", "Allos", "ALLOS ON"],
    "FS Bio": ["FS Bio", "FS Bioenergia"],
    "Sports Media": ["Sports Media", "Sports Media Futebol Brasileiro", "Sports Media SA", "RBRAB7"],
    "Alares": ["Alares", "CONX12", "Triple Play Brasil Participações", "Triple Play Brasil Participações S.A."],
    "Electra Energy": ["Electra Energy", "FIDC Siga", "Electra Energia", "Grupo Electra Energia"],
    "JSLG3": ["JSL", "JSL SA", "JSLG3", "Grupo JSL"],
    "TOTS3": ["TOTVS", "TOTS3", "TOTVS S.A"],
    "MDNE3": ["Moura Dubeux", "Moura Dubeux Engenharia", "Moura Dubeux Engenharia SA", "MDNE3"],
    "AXIA3": ["AXIA", "AXIA Energia", "AXIA3", "AXIA6"],
    "KLBN11": ["Klabin", "KLBN11"],
    "CSAN3": ["Cosan", "Cosan S.A.", "CSAN3"],
    "Solví": ["Solví", "Solví Energia Verde"],
}

COLUNAS = ["Data", "Hora", "Emissor", "Título", "Fonte", "Link"]

# ============================================================
# FUNÇÕES AUXILIARES
# ============================================================


def carregar_historico_clipping():
    if os.path.exists(HISTORICO_JSON):
        with open(HISTORICO_JSON, "r", encoding="utf-8") as f:
            return set(json.load(f))
    return set()


def salvar_historico_clipping(historico):
    with open(HISTORICO_JSON, "w", encoding="utf-8") as f:
        json.dump(sorted(historico), f, ensure_ascii=False, indent=2)


def normalizar_titulo(titulo):
    if not isinstance(titulo, str):
        return ""
    titulo = unicodedata.normalize("NFKD", titulo)
    titulo = "".join(c for c in titulo if not unicodedata.combining(c))
    titulo = titulo.lower()
    titulo = re.sub(r"[^a-z0-9\s]", "", titulo)
    stopwords = [
        "diz", "afirma", "veja", "saiba", "entenda",
        "empresa", "empresas", "mercado", "brasil",
        "hoje", "agora", "segundo", "aponta",
    ]
    palavras = [p for p in titulo.split() if p not in stopwords]
    return " ".join(palavras)


def link_valido(url):
    try:
        r = requests.head(url, allow_redirects=True, timeout=5)
        return r.status_code == 200
    except Exception:
        return False


def encurtar_link(url):
    if not link_valido(url):
        return None
    try:
        r = requests.get(
            "https://tinyurl.com/api-create.php",
            params={"url": url},
            timeout=5,
        )
        if r.status_code == 200 and r.text.startswith("http"):
            return r.text
    except Exception:
        pass
    return None


def noticia_relevante(entry, termos):
    def norm(t):
        t = unicodedata.normalize("NFKD", t)
        t = "".join(c for c in t if not unicodedata.combining(c))
        return t.lower()

    texto = norm(entry.title) + " " + norm(getattr(entry, "summary", ""))
    return any(norm(termo) in texto for termo in termos)


def buscar_rss(query, hoje):
    rss_url = (
        "https://news.google.com/rss/search?"
        f"q={urllib.parse.quote(query)}&hl=pt-BR&gl=BR&ceid=BR:pt"
    )
    feed = feedparser.parse(rss_url)
    entradas = []
    for entry in feed.entries:
        if not getattr(entry, "published_parsed", None):
            continue
        # O feed vem em UTC; converte para o horário de Brasília antes de comparar
        data_pub = (
            datetime(*entry.published_parsed[:6], tzinfo=timezone.utc)
            .astimezone(FUSO)
            .date()
        )
        if data_pub != hoje:
            continue
        entradas.append(entry)
    return entradas


MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]

# Paleta do modelo de e-mail da Tyton
COR_VERDE = "#2a3e32"
COR_LINHA = "#c9d3cd"
COR_ZEBRA = "#f2f4f3"
COR_FUNDO = "#eef1ef"
FONTE = "'Segoe UI',Calibri,Arial,sans-serif"


def montar_html(df, agora):
    data_cab = f"{agora.day:02d} {MESES[agora.month - 1]} {agora.year} &nbsp;·&nbsp; {agora:%Hh%M}"
    n_emissores = df["Emissor"].nunique()
    plural = "s" if len(df) != 1 else ""

    secoes = ""
    for i, (emissor, grupo) in enumerate(df.groupby("Emissor"), start=1):
        linhas = ""
        for j, (_, row) in enumerate(grupo.iterrows()):
            fundo = COR_ZEBRA if j % 2 == 0 else "#ffffff"
            # O Google News repete a fonte no fim do título; ela já aparece no link abaixo
            titulo = row["Título"]
            sufixo = f" - {row['Fonte']}"
            if titulo.endswith(sufixo):
                titulo = titulo[: -len(sufixo)]
            linhas += f"""
<tr>
  <td style="background:{fundo}; padding:10px 14px; border-bottom:1px solid #e3e8e5;">
    <div style="font-size:10.5pt; font-weight:600; color:#222; line-height:1.45;">{html_lib.escape(titulo)}</div>
    <a href="{row["Link"]}" style="font-size:9pt; color:{COR_VERDE}; text-decoration:underline;">{html_lib.escape(row["Fonte"])}</a>
  </td>
</tr>"""
        secoes += f"""
<tr><td style="padding:22px 32px 0 32px;">
  <div style="font-size:8.5pt; font-weight:700; letter-spacing:1.5px; color:#333;
              padding-bottom:6px; border-bottom:1px solid {COR_LINHA};">
    {i:02d} &nbsp;·&nbsp; {html_lib.escape(emissor).upper()}
  </div>
</td></tr>
<tr><td style="padding:8px 16px 0 16px;">
  <table width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse;">{linhas}
  </table>
</td></tr>"""

    return f"""<html>
<body style="margin:0; padding:0; background:{COR_FUNDO}; font-family:{FONTE}; color:#222;">
<table width="100%" cellpadding="0" cellspacing="0" style="background:{COR_FUNDO};">
<tr><td align="center" style="padding:24px 12px;">
<table width="640" cellpadding="0" cellspacing="0" style="max-width:640px; width:100%; background:#ffffff; border-collapse:collapse;">

<tr><td style="background:{COR_VERDE}; padding:24px 32px;">
  <table width="100%" cellpadding="0" cellspacing="0"><tr>
    <td style="color:#ffffff; font-size:13pt; font-weight:700; letter-spacing:4px;">TYTON CAPITAL</td>
    <td align="right" style="color:#d9e2dc; font-size:8.5pt; letter-spacing:1px;">CLIPPING DIÁRIO &nbsp;|&nbsp; {data_cab}</td>
  </tr></table>
</td></tr>

<tr><td style="padding:28px 32px 0 32px; font-size:11pt; line-height:1.6;">
  <p style="margin:0 0 12px 0;">Prezados,</p>
  <p style="margin:0;">Segue o clipping com as notícias publicadas hoje sobre os emissores da carteira:
  <b>{len(df)} notícia{plural}</b> de <b>{n_emissores} emissor{"es" if n_emissores != 1 else ""}</b>.</p>
</td></tr>
{secoes}

<tr><td style="padding:24px 32px 28px 32px;">
  <div style="font-size:8pt; color:#777; border-top:1px solid {COR_LINHA}; padding-top:10px; line-height:1.5;">
    Notícias do Google News publicadas em {agora:%d/%m/%Y}. Cada envio traz só o que ainda não foi enviado.
    A planilha com todas as notícias do dia segue em anexo.<br>
    Gerado automaticamente · Tyton Capital
  </div>
</td></tr>

</table>
</td></tr>
</table>
</body></html>
"""


def salvar_excel(df_novas, hoje_str):
    # Acumula as notícias do dia; quando o dia vira, começa uma planilha nova
    if os.path.exists(CAMINHO_EXCEL):
        df_existente = pd.read_excel(CAMINHO_EXCEL)
        df_existente = df_existente[df_existente["Data"] == hoje_str]
        df_dia = pd.concat([df_existente, df_novas], ignore_index=True)
    else:
        df_dia = df_novas
    df_dia = df_dia[COLUNAS]
    df_dia.to_excel(CAMINHO_EXCEL, index=False)

    wb = load_workbook(CAMINHO_EXCEL)
    ws = wb.active
    tabela = Table(displayName="Noticias_diarias", ref=f"A1:F{ws.max_row}")
    tabela.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tabela)
    wb.save(CAMINHO_EXCEL)


def obter_token_graph():
    r = requests.post(
        f"https://login.microsoftonline.com/{GRAPH_TENANT_ID}/oauth2/v2.0/token",
        data={
            "client_id": GRAPH_CLIENT_ID,
            "client_secret": GRAPH_CLIENT_SECRET,
            "scope": "https://graph.microsoft.com/.default",
            "grant_type": "client_credentials",
        },
        timeout=30,
    )
    r.raise_for_status()
    return r.json()["access_token"]


def enviar_email(html, agora):
    with open(CAMINHO_EXCEL, "rb") as f:
        anexo_b64 = base64.b64encode(f.read()).decode()

    mensagem = {
        "message": {
            "subject": f"Tyton: Clipping {agora.strftime('%d/%m/%Y')} ({agora.strftime('%Hh%M')})",
            "body": {"contentType": "HTML", "content": html},
            "toRecipients": [
                {"emailAddress": {"address": e.strip()}}
                for e in EMAIL_DESTINO.split(",") if e.strip()
            ],
            "attachments": [{
                "@odata.type": "#microsoft.graph.fileAttachment",
                "name": "Noticias_Diarias.xlsx",
                "contentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "contentBytes": anexo_b64,
            }],
        },
        "saveToSentItems": True,
    }
    r = requests.post(
        f"https://graph.microsoft.com/v1.0/users/{EMAIL_REMETENTE}/sendMail",
        headers={"Authorization": f"Bearer {obter_token_graph()}"},
        json=mensagem,
        timeout=60,
    )
    if r.status_code != 202:
        raise RuntimeError(f"Graph recusou o envio ({r.status_code}): {r.text}")


# ============================================================
# EXECUÇÃO
# ============================================================


def main():
    # --previa: coleta e gera o HTML em dados/previa.html, sem enviar nem gravar histórico
    previa = "--previa" in sys.argv
    faltando = [
        v for v in ("CLIPPING_TENANT_ID", "CLIPPING_CLIENT_ID", "CLIPPING_CLIENT_SECRET")
        if not os.environ.get(v)
    ]
    if faltando and not previa:
        sys.exit(f"❌ Variáveis de ambiente não definidas: {', '.join(faltando)}")

    os.makedirs(PASTA_DADOS, exist_ok=True)
    agora = datetime.now(FUSO)
    hoje = agora.date()
    hoje_str = hoje.strftime("%d/%m/%Y")

    # ---------- COLETA (query geral + query por sites) ----------
    historico_clipping = carregar_historico_clipping()
    noticias = []
    for emissor, termos in EMISSORES.items():
        base_query = " OR ".join([f'"{t}"' for t in termos])
        entradas_a = buscar_rss(base_query, hoje)
        entradas_b = buscar_rss(f"({base_query}) ({SITES_FINANCEIROS})", hoje)

        for entry in entradas_a + entradas_b:
            if not noticia_relevante(entry, termos):
                continue
            # Evita validar e encurtar link de notícia que já foi enviada
            if normalizar_titulo(entry.title) in historico_clipping:
                continue
            link_curto = encurtar_link(entry.link)
            if not link_curto:
                continue
            fonte = entry.source.title if "source" in entry else "Google News"
            noticias.append({
                "Data": hoje_str,
                "Hora": agora.strftime("%H:%M"),
                "Emissor": emissor,
                "Título": entry.title,
                "Fonte": fonte,
                "Link": link_curto,
            })

    # ---------- ANTI-REPETIÇÃO (dentro do clipping + entre clippings) ----------
    df = pd.DataFrame(noticias, columns=COLUNAS)
    df["titulo_normalizado"] = df["Título"].apply(normalizar_titulo)
    df = df.drop_duplicates(subset=["Emissor", "titulo_normalizado"])
    df = df.sort_values(["Emissor", "Título"])
    df = df.drop_duplicates(subset=["titulo_normalizado"], keep="first")
    df = df[~df["titulo_normalizado"].isin(historico_clipping)]
    df = df.groupby("Emissor").head(6).reset_index(drop=True)

    if df.empty:
        print(f"ℹ️ {agora:%d/%m %H:%M}: nenhuma notícia nova. Email não enviado.")
        return

    novos_titulos = set(df["titulo_normalizado"])
    df = df.drop(columns=["titulo_normalizado"])

    if previa:
        with open(os.path.join(PASTA_DADOS, "previa.html"), "w", encoding="utf-8") as f:
            f.write(montar_html(df, agora))
        print(f"👀 Prévia com {len(df)} notícias salva em dados/previa.html (nada enviado).")
        return

    # ---------- EXCEL + EMAIL ----------
    salvar_excel(df, hoje_str)
    enviar_email(montar_html(df, agora), agora)

    # Só grava no histórico depois que o email saiu
    historico_clipping.update(novos_titulos)
    salvar_historico_clipping(historico_clipping)

    print(f"✅ {agora:%d/%m %H:%M}: {len(df)} notícias novas enviadas para {EMAIL_DESTINO}.")


if __name__ == "__main__":
    main()
